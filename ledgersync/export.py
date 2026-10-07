"""A client's work as an Excel workbook (design of 2026-10-06): its trial balance and its transactions, from
the same ledger the pages show. Amounts are numbers in pounds and dates are dates, so the sheets add up."""
from __future__ import annotations

import datetime as dt
import io
import re
from typing import Union

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font

from .models import Direction, Ledger, TrialBalance

MONEY = '"£"#,##0.00'
DATE = "dd/mm/yyyy"
COLUMNS = ("Date", "Description", "Counterparty", "Document type", "In/Out", "Amount", "VAT", "Net", "Account",
           "Account name", "Other side", "Still owed", "Issues", "Upload", "Edited")
_TYPE_NAMES = {"receipt": "Receipt", "invoice": "Invoice", "expense_claim": "Expense claim",
               "statement": "Bank statement", "quote": "Quote", "pro_forma": "Pro forma",
               "purchase_order": "Purchase order", "remittance_advice": "Remittance advice",
               "supplier_statement": "Supplier statement", "other": "Other"}
_UNSAFE = re.compile(r'[\\/:*?"<>|\x00-\x1f]')
_SPACES = re.compile(r"\s+")


def file_name(client_name: str, day: dt.date) -> str:
    """'Business Cube Ltd 2026-10-06.xlsx', without characters that file names can't hold."""
    name = _SPACES.sub(" ", _UNSAFE.sub(" ", client_name)).strip(" .") or "Client"
    return f"{name} {day.isoformat()}.xlsx"


def workbook(ledger: Ledger, balance: Union[TrialBalance, list[str]], day: dt.date) -> bytes:
    """The workbook: a "Trial balance" sheet (or, while rows need fixing, what stops it) and a "Transactions"
    sheet with one row per saved row, in the table's order."""
    book = Workbook()
    _trial_balance(book.active, ledger, balance, day)
    _transactions(book.create_sheet("Transactions"), ledger)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def _append(sheet, values) -> None:
    """Adds a row. Text read from documents stays text: characters a worksheet can't hold are dropped, and
    text starting with "=" is not made a formula."""
    sheet.append([ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v for v in values])
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str) and cell.value.startswith("="):
            cell.data_type = "s"


def _title(sheet, text: str) -> None:
    _append(sheet, [text])
    sheet["A1"].font = Font(bold=True)


def _widths(sheet, widths) -> None:
    for n, width in enumerate(widths):
        sheet.column_dimensions[chr(ord("A") + n)].width = width


def _trial_balance(sheet, ledger: Ledger, balance, day: dt.date) -> None:
    sheet.title = "Trial balance"
    _title(sheet, f"{ledger.client.name} — trial balance — {day.day} {day:%b %Y}")
    if isinstance(balance, list):
        _append(sheet, ["The trial balance can't be produced yet:"])
        for problem in balance:
            _append(sheet, [problem])
        return
    _append(sheet, ["Code", "Account", "Debit", "Credit"])
    for line in balance.lines:
        _append(sheet, [line.code, line.name, line.debit, line.credit])
    _append(sheet, [None, "Totals", balance.total_debits, balance.total_credits])
    for row in sheet.iter_rows(min_row=3, min_col=3, max_col=4):
        for cell in row:
            cell.number_format = MONEY
    sheet.freeze_panes = "A3"
    _widths(sheet, (8, 34, 14, 14))


def _transactions(sheet, ledger: Ledger) -> None:
    uploads = {u.id: u.name for u in ledger.uploads}
    _title(sheet, f"{ledger.client.name} — transactions")
    _append(sheet, COLUMNS)
    for tx in ledger.transactions:
        code, name = (tx.paid_against, tx.paid_against_name) if tx.paid_against else (tx.account_code, tx.account_name)
        _append(sheet, [tx.date, tx.description, tx.counterparty, _TYPE_NAMES.get(tx.document_type, tx.document_type),
                      "In" if tx.direction == Direction.IN else "Out", tx.gross, tx.vat_posted, tx.net, code, name,
                      tx.contra_account_code, tx.owed, "; ".join(i.message for i in tx.issues) or None,
                      uploads.get(tx.upload_id), "Yes" if tx.edited else None])
    for row in sheet.iter_rows(min_row=3):
        row[0].number_format = DATE
        for cell in (*row[5:8], row[11]):
            cell.number_format = MONEY
    sheet.freeze_panes = "A3"
    _widths(sheet, (12, 36, 22, 16, 8, 12, 10, 12, 9, 26, 10, 12, 50, 24, 8))
